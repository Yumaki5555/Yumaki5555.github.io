# 総務省「全国地方公共団体コード」の Excel から、地域選択用の市区町村一覧（data/municipalities.json）を作ります。
#
# 使い方：python _tools/build_municipalities.py <ダウンロードした xlsx のパス> <更新日の表記>
#   例）python _tools/build_municipalities.py code.xlsx "令和6年1月1日更新"
#   Excel は https://www.soumu.go.jp/denshijiti/code.html の「都道府県コード及び市区町村コード」
#
# コードは6けた（5けたの団体コード＋検査数字）。政令指定都市は市全体で1件とし、区は入れません。
# 北方領土の6村（01695〜01700）は、どの専門サイトにもページがないので入れません
# （全国 多文化共生予算マップの市区町村一覧と同じ範囲：47都道府県＋1,741市区町村）。
import json
import sys
import unicodedata
from pathlib import Path

import openpyxl


def hira(s):
    s = unicodedata.normalize("NFKC", s or "").replace(" ", "").replace("　", "")
    return "".join(chr(ord(c) - 0x60) if "ァ" <= c <= "ヶ" else c for c in s)


def main():
    src, ver = Path(sys.argv[1]), sys.argv[2]
    out = Path(__file__).resolve().parent.parent / "data" / "municipalities.json"
    wb = openpyxl.load_workbook(src, read_only=True)
    # 2枚目のシート（政令指定都市とその区）にある市のコード
    rows2 = list(wb.worksheets[1].iter_rows(values_only=True))[1:]
    seirei_codes = {str(r[0]).zfill(6) for r in rows2 if r[0]}
    prefs, munis = [], []
    for row in list(wb.worksheets[0].iter_rows(values_only=True))[1:]:
        code, pname, mname, pkana, mkana = (str(row[0]).zfill(6),) + tuple(row[1:5])
        if not pname:
            continue
        if mname:
            if "01695" <= code[:5] <= "01700":
                continue
            m = [code, mname, hira(mkana)]
            if code in seirei_codes:
                m.append(1)  # 政令指定都市
            munis.append(m)
        else:
            prefs.append([code[:2], pname, hira(pkana)])
    data = {
        "source": "総務省「全国地方公共団体コード」（" + ver + "）",
        "url": "https://www.soumu.go.jp/denshijiti/code.html",
        "prefs": prefs,
        "munis": munis,
    }
    text = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    out.write_text(text, encoding="utf-8", newline="\n")
    print(out.name, len(prefs), "prefs", len(munis), "munis", len(text.encode("utf-8")), "bytes")


if __name__ == "__main__":
    main()
